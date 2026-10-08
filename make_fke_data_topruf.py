"""Exporta el caso de Papelería del seed a cuatro Excel, sin conectarse a la BD."""
from datetime import date,datetime
import argparse
from decimal import Decimal
from importlib.util import module_from_spec,spec_from_file_location
from pathlib import Path
import sys
import pandas as pd
RAIZ=Path(__file__).resolve().parent; 
BACKEND=RAIZ/"backend"; 
INICIO=date(2026,1,1); 
FIN=date(2026,9,30)
# Cambia estos valores o usa los argumentos de línea de comandos.
PARAMETROS={"cantidad_productos":35,
            "cantidad_ventas":4463,
            "cantidad_compras":329,
            "cantidad_gastos":290,
            "ruta_salida":RAIZ/"datos_prueba","semilla":2026}
def seleccionar(filas,cantidad,nombre):
 if cantidad is None or cantidad==len(filas): return list(filas)
 if cantidad<0 or cantidad>len(filas): raise ValueError("Solicitaste "+str(cantidad)+" "+nombre+", disponibles: "+str(len(filas)))
 if cantidad==0: return []
 indices=sorted({round(i*(len(filas)-1)/(cantidad-1)) if cantidad>1 else len(filas)//2 for i in range(cantidad)})
 return [filas[i] for i in indices]
def seed():
 sys.path.insert(0,str(BACKEND)); spec=spec_from_file_location("seed_cc",BACKEND/"scripts"/"generar_seed.py")
 mod=module_from_spec(spec); sys.modules[spec.name]=mod; spec.loader.exec_module(mod); return mod

def generar(filas,parametros):
 ide=next(e["id_empresa"] for e in filas["empresas"] if e["giro"]=="Papelería")
 ps=sorted([p for p in filas["productos_cat"] if p["id_empresa"]==ide],key=lambda x:x["id_producto"])
 nprod=parametros["cantidad_productos"]
 if not 30<=nprod<=len(ps): raise ValueError("cantidad_productos debe estar entre 30 y "+str(len(ps)))
 ps=ps[:nprod]; nombres={p["id_producto"] for p in ps}; pi={p["id_producto"]:p for p in ps}
 vs=sorted([x for x in filas["historial_ventas"] if x["id_empresa"]==ide and INICIO<=x["fecha_hora"].date()<=FIN and x["id_producto"] in nombres],key=lambda x:x["fecha_hora"])
 cs=sorted([x for x in filas["compras_producto"] if x["id_empresa"]==ide and INICIO<=x["fecha"]<=FIN and x["id_producto"] in nombres],key=lambda x:x["fecha"])
 gs=sorted([x for x in filas["gastos_operativos"] if x["id_empresa"]==ide and INICIO<=x["fecha"]<=FIN],key=lambda x:x["fecha"])
 vs=seleccionar(vs,parametros["cantidad_ventas"],"ventas"); cs=seleccionar(cs,parametros["cantidad_compras"],"compras"); gs=seleccionar(gs,parametros["cantidad_gastos"],"gastos")
 prod=[]
 for p in ps:
  mov=[]
  mov.extend((x["fecha"],int(x["cantidad"])) for x in cs if x["id_producto"]==p["id_producto"])
  mov.extend((x["fecha_hora"],-int(x["cantidad_vendida"])) for x in vs if x["id_producto"]==p["id_producto"])
  saldo=0; minimo=0
  for _,delta in sorted(mov,key=lambda x:x[0].date() if isinstance(x[0],datetime) else x[0]): saldo+=delta; minimo=min(minimo,saldo)
  existencia=-minimo+saldo
  prod.append({"Producto":p["sku_o_nombre"],"Categoría":p["categoria"],"Precio de venta":p["precio_venta"],"Costo":p["costo_promedio"],"Existencia":existencia,"Mínimo":p["stock_minimo"],"Unidad":p["unidad"]})
 ventas=[]
 for x in vs:
  p=pi[x["id_producto"]]; ventas.append({"Fecha":x["fecha_hora"],"Producto":p["sku_o_nombre"],"Cantidad":x["cantidad_vendida"],"Precio unitario":x["precio_unitario"],"Costo unitario":x["costo_unitario"],"Total de la venta":x["cantidad_vendida"]*x["precio_unitario"],"Categoría":p["categoria"]})
 compras=[]
 for x in cs:
  p=pi[x["id_producto"]]; compras.append({"Fecha":x["fecha"],"Producto":p["sku_o_nombre"],"Cantidad":x["cantidad"],"Costo unitario":x["costo_unitario"],"Total":x["cantidad"]*x["costo_unitario"],"Proveedor":x["proveedor"] or "Papelera del Centro"})
 gastos=[{"Fecha":x["fecha"],"Concepto":x["concepto"],"Monto":x["monto"],"Categoría":x["categoria"],"Tipo":x["tipo"]} for x in gs]
 cols={"productos_prueba.xlsx":["Producto","Categoría","Precio de venta","Costo","Existencia","Mínimo","Unidad"],"ventas_prueba.xlsx":["Fecha","Producto","Cantidad","Precio unitario","Costo unitario","Total de la venta","Categoría"],"compras_prueba.xlsx":["Fecha","Producto","Cantidad","Costo unitario","Total","Proveedor"],"gastos_prueba.xlsx":["Fecha","Concepto","Monto","Categoría","Tipo"]}
 raw={"productos_prueba.xlsx":prod,"ventas_prueba.xlsx":ventas,"compras_prueba.xlsx":compras,"gastos_prueba.xlsx":gastos}; return {n:pd.DataFrame(raw[n],columns=c) for n,c in cols.items()}

def escribir(df,path,tipo):
 import io
 from openpyxl import load_workbook
 from app.ingestion.plantillas import generar as generar_plantilla
 wb=load_workbook(io.BytesIO(generar_plantilla(tipo)))
 ws=wb.worksheets[0]
 # Conserva Instrucciones; sustituye solo filas de ejemplo.
 if ws.max_row>1: ws.delete_rows(2,ws.max_row-1)
 for i,row in enumerate(df.itertuples(index=False,name=None),start=2):
  for j,value in enumerate(row,start=1):
   if pd.isna(value): value=None
   cell=ws.cell(i,j,value)
   if isinstance(value,(date,datetime)): cell.number_format="dd/mm/yyyy"
 wb.save(path)

def validar(d,folder):
 p,v,c,g=(d["productos_prueba.xlsx"],d["ventas_prueba.xlsx"],d["compras_prueba.xlsx"],d["gastos_prueba.xlsx"])
 assert len(p)>=30 and len(v)>=2000 and not p.Producto.duplicated().any() and not v.duplicated().any() and not c.duplicated().any() and not g.duplicated().any()
 assert set(v.Producto)==set(c.Producto)==set(p.Producto) and v.Fecha.dt.date.between(INICIO,FIN).all(), (set(p.Producto)-set(v.Producto),set(p.Producto)-set(c.Producto),v.Fecha.min(),v.Fecha.max())
 assert all(t==q*z for t,q,z in zip(v["Total de la venta"],v.Cantidad,v["Precio unitario"])) and all(t==q*z for t,q,z in zip(c.Total,c.Cantidad,c["Costo unitario"]))
 ventas=sum((Decimal(str(x)) for x in v["Total de la venta"]),Decimal(0)); costo=sum((Decimal(str(q))*Decimal(str(z)) for q,z in zip(v.Cantidad,v["Costo unitario"])),Decimal(0)); gas=sum((Decimal(str(x)) for x,t in zip(g.Monto,g.Tipo) if t!="retiro"),Decimal(0))

 vend=v.groupby("Producto").Cantidad.sum().to_dict(); comp=c.groupby("Producto").Cantidad.sum().to_dict()
 for _,x in p.iterrows(): assert Decimal(str(x.Existencia))+Decimal(str(vend[x.Producto]))-Decimal(str(comp[x.Producto]))>=0
 sys.path.insert(0,str(BACKEND))
 from app.ingestion.lectura import leer_archivo
 from app.ingestion.mapeo import por_reglas,faltantes
 from app.ingestion.limpieza import limpiar
 tipos={"productos_prueba.xlsx":"productos","ventas_prueba.xlsx":"ventas","compras_prueba.xlsx":"compras","gastos_prueba.xlsx":"gastos"}
 for n,t in tipos.items():
  tabla=leer_archivo(n,(folder/n).read_bytes(),5); m=por_reglas(t,tabla.columnas); assert not faltantes(t,m); r=limpiar(t,tabla.columnas,tabla.filas,m,tabla.numeros,hoy=FIN); assert len(r.filas)==len(d[n]) and not r.errores,(t,r.errores[:2])
 return ventas,costo,gas

def main():
 parser=argparse.ArgumentParser(description="Generador Excel de Papelería El Lápiz Feliz")
 parser.add_argument("--productos",type=int,default=PARAMETROS["cantidad_productos"])
 parser.add_argument("--ventas",type=int,default=PARAMETROS["cantidad_ventas"])
 parser.add_argument("--compras",type=int,default=PARAMETROS["cantidad_compras"])
 parser.add_argument("--gastos",type=int,default=PARAMETROS["cantidad_gastos"],help="Total de gastos y retiros")
 parser.add_argument("--salida",default=str(PARAMETROS["ruta_salida"]))
 parser.add_argument("--semilla",type=int,default=PARAMETROS["semilla"])
 args=parser.parse_args(); opts={"cantidad_productos":args.productos,"cantidad_ventas":args.ventas,"cantidad_compras":args.compras,"cantidad_gastos":args.gastos}
 modulo=seed(); modulo.SEMILLA=args.semilla; d=generar(modulo.construir(),opts); folder=Path(args.salida).expanduser().resolve(); folder.mkdir(parents=True,exist_ok=True); sheets={"productos_prueba.xlsx":"productos","ventas_prueba.xlsx":"ventas","compras_prueba.xlsx":"compras","gastos_prueba.xlsx":"gastos"}
 for n,df in d.items(): escribir(df,folder/n,sheets[n])
 a,b,c=validar(d,folder); print("Generados y validados en "+str(folder))
 for n,df in d.items(): print(n+": "+format(len(df),",")+" filas")
 print("Ene-sep 2026: ventas $"+format(a,",.2f")+"; costo $"+format(b,",.2f")+"; gastos operativos $"+format(c,",.2f"))
if __name__=="__main__": main()
