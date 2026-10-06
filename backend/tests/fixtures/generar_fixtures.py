"""Genera los archivos de prueba de ingesta (formatos realistas de tienditas).

    python tests/fixtures/generar_fixtures.py

1. papeleria_ventas.xlsx   Título arriba, fechas de Excel, "P. Unitario", fila de TOTAL, un precio "N/A".
2. abarrotes_ventas.csv    Separador ';', acentos en Windows-1252, coma decimal, fechas dd/mm/aaaa.
3. boutique_compras.xlsx   Montos como texto "$1,250.00", sin costo unitario (solo Total), una fila repetida.
4. gastos_tienda.csv       UTF-8 con BOM, fechas ISO, sin columna de tipo, una celda con fórmula maliciosa.
5. catalogo_productos.xlsx Encabezados "Descripción / Existencia / Precio público / Costo".
"""

import csv
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

AQUI = Path(__file__).parent


def papeleria() -> None:
    libro = Workbook()
    h = libro.active
    h.append(["PAPELERÍA EL LÁPIZ FELIZ"])
    h.append(["Reporte de ventas - septiembre 2026"])
    h.append([])
    h.append(["Fecha", "Artículo", "Cant.", "P. Unitario", "Importe"])
    filas = [
        (datetime(2026, 9, 1, 9, 30), "Cuaderno profesional raya 100 h", 3, 50, 150),
        (datetime(2026, 9, 1, 11, 5), "Pluma BIC azul", 10, 8, 80),
        (datetime(2026, 9, 2, 16, 40), "Copias blanco y negro", 120, 1, 120),
        (datetime(2026, 9, 3, 10, 0), "Producto nuevo de prueba", 2, 35, 70),
        (datetime(2026, 9, 3, 12, 0), "Colores Maped 12", 1, "N/A", 0),
    ]
    for f in filas:
        h.append(list(f))
    h.append([])
    h.append(["TOTAL", "", 136, "", 420])
    libro.save(AQUI / "papeleria_ventas.xlsx")


def abarrotes() -> None:
    with open(AQUI / "abarrotes_ventas.csv", "w", encoding="cp1252", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["FECHA", "PRODUCTO", "PIEZAS", "PRECIO VENTA", "COSTO"])
        w.writerow(["15/09/2026", "Leche Lala 1 L", "4", "30,50", "26,10"])
        w.writerow(["15/09/2026", "Huevo blanco", "1,5", "53,50", "48,05"])
        w.writerow(["16/09/2026", "Azúcar", "2", "36", "28,40"])
        w.writerow(["31/02/2026", "Arroz", "1", "34", "25"])          # fecha imposible
        w.writerow(["17/09/2026", "", "1", "20", "15"])                # sin producto


def boutique() -> None:
    libro = Workbook()
    h = libro.active
    h.append(["Fecha de compra", "Prenda", "Proveedor", "Unidades", "Total"])
    h.append(["05/09/2026", "Chamarra acolchada", "Importadora Angelópolis", 10, "$5,600.00"])
    h.append(["05/09/2026", "Suéter tejido", "Tejidos Chignahuapan", 12, "$3,000.00"])
    h.append(["05/09/2026", "Suéter tejido", "Tejidos Chignahuapan", 12, "$3,000.00"])   # repetida
    h.append(["06/09/2026", "Bufanda", "Tejidos Chignahuapan", 0, "$0.00"])             # cantidad 0
    libro.save(AQUI / "boutique_compras.xlsx")


def gastos() -> None:
    with open(AQUI / "gastos_tienda.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["dia", "descripcion", "monto"])
        w.writerow(["2026-09-01", "Renta del local", "$1,200"])
        w.writerow(["2026-09-05", "Bolsas y empaque", "85.50"])
        w.writerow(["2026-09-10", "Internet Telmex", "389"])
        w.writerow(["2026-09-12", '=HYPERLINK("http://malicioso.example","clic")', "50"])
        w.writerow(["2026-09-30", "Retiro personal", "3000"])
        w.writerow(["2026-09-30", "Gasolina", "-20"])                 # negativo


def catalogo() -> None:
    libro = Workbook()
    h = libro.active
    h.append(["Código", "Descripción", "Existencia", "Precio público", "Costo", "Línea"])
    h.append(["A-01", "Cuaderno profesional raya 100 h", 40, 50, 22.12, "Cuadernos"])
    h.append(["A-02", "Producto nuevo de catálogo", 15, 99.9, 60, "Varios"])
    libro.save(AQUI / "catalogo_productos.xlsx")


if __name__ == "__main__":
    for generar in (papeleria, abarrotes, boutique, gastos, catalogo):
        generar()
    print("Fixtures generados en", AQUI)
