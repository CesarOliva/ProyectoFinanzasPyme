"""Lectura segura y profiling de archivos .xlsx / .csv.

- Valida extensión, tamaño y firma del archivo.
- Lee todo como texto: nunca evalúa fórmulas (openpyxl data_only=True usa el valor guardado).
- Detecta la fila de encabezados aunque haya títulos arriba, y quita filas vacías o de totales.
"""

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, time

from openpyxl import load_workbook


class ArchivoInvalido(ValueError):
    """Error con mensaje listo para mostrar al usuario."""


@dataclass
class TablaLeida:
    columnas: list[str]
    filas: list[list[str]]
    fila_encabezado: int                 # número de fila (1 = primera) en el archivo original
    numeros: list[int] = field(default_factory=list)   # número de fila original de cada fila de datos
    advertencias: list[str] = field(default_factory=list)


def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m-%d %H:%M:%S") if valor.time() != time(0, 0) else valor.strftime("%Y-%m-%d")
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _leer_xlsx(contenido: bytes) -> list[list[str]]:
    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as error:  # openpyxl lanza varios tipos según el daño del archivo
        raise ArchivoInvalido("No pude abrir el Excel. Verifica que sea un .xlsx válido y no esté protegido.") from error
    hoja = libro.worksheets[0]
    filas = [[_texto(c) for c in fila] for fila in hoja.iter_rows(values_only=True)]
    libro.close()
    return filas


def _leer_csv(contenido: bytes) -> list[list[str]]:
    for codificacion in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = contenido.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ArchivoInvalido("No pude leer el texto del CSV.")
    muestra = texto[:5000]
    try:
        delimitador = csv.Sniffer().sniff(muestra, delimiters=",;\t|").delimiter
    except csv.Error:
        delimitador = ";" if muestra.count(";") > muestra.count(",") else ","
    return [[c.strip() for c in fila] for fila in csv.reader(io.StringIO(texto), delimiter=delimitador)]


def _es_numero(texto: str) -> bool:
    return bool(re.fullmatch(r"[\s$\-+(]*[\d.,]+\)?\s*(mxn)?", texto.lower()))


def _es_fila_total(fila: list[str]) -> bool:
    primeras = " ".join(c.lower() for c in fila[:3] if c)
    return bool(re.match(r"^\s*(gran\s+)?total(es)?\b", primeras))


def _detectar_encabezado(filas: list[list[str]]) -> int:
    """Índice de la fila más parecida a un encabezado dentro de las primeras 20."""
    mejor, puntaje_mejor = 0, -1.0
    for i, fila in enumerate(filas[:20]):
        llenas = [c for c in fila if c]
        if len(llenas) < 2:
            continue
        textos = sum(not _es_numero(c) for c in llenas)
        siguientes = [f for f in filas[i + 1:i + 4] if any(f)]
        con_datos = sum(1 for f in siguientes if sum(bool(c) for c in f) >= len(llenas) * 0.5)
        puntaje = textos / len(llenas) * 2 + len(llenas) * 0.1 + con_datos
        if puntaje > puntaje_mejor:
            mejor, puntaje_mejor = i, puntaje
    return mejor


def normalizar_encabezado(texto: str) -> str:
    sin_acentos = "".join(c for c in unicodedata.normalize("NFD", texto.lower()) if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", sin_acentos)).strip()


def leer_archivo(nombre: str, contenido: bytes, max_mb: int) -> TablaLeida:
    extension = nombre.lower().rsplit(".", 1)[-1] if "." in nombre else ""
    if extension not in ("xlsx", "csv"):
        raise ArchivoInvalido("Solo acepto archivos .xlsx (Excel) o .csv.")
    if len(contenido) > max_mb * 1024 * 1024:
        raise ArchivoInvalido(f"El archivo pesa más de {max_mb} MB. Divídelo en partes más pequeñas.")
    if not contenido:
        raise ArchivoInvalido("El archivo está vacío.")
    if extension == "xlsx" and not contenido.startswith(b"PK"):
        raise ArchivoInvalido("El archivo dice ser Excel pero su contenido no lo es.")

    crudas = _leer_xlsx(contenido) if extension == "xlsx" else _leer_csv(contenido)
    # Se conserva el número de fila original para reportar errores como los ve el usuario en Excel.
    numeradas = [(n, f) for n, f in enumerate(crudas, start=1) if any(c for c in f)]
    filas = [f for _, f in numeradas]
    if not filas:
        raise ArchivoInvalido("No encontré datos en el archivo.")

    i = _detectar_encabezado(filas)
    encabezado = filas[i]
    ancho = max(len(f) for f in filas)
    encabezado = encabezado + [""] * (ancho - len(encabezado))
    advertencias = []
    if i > 0:
        advertencias.append(f"Ignoré {i} fila(s) de título antes de los encabezados.")

    datos, numeros, totales = [], [], 0
    for numero, fila in numeradas[i + 1:]:
        fila = fila + [""] * (ancho - len(fila))
        if _es_fila_total(fila):
            totales += 1
            continue
        datos.append(fila)
        numeros.append(numero)
    if totales:
        advertencias.append(f"Ignoré {totales} fila(s) de totales.")

    # Quita columnas sin encabezado y sin datos.
    utiles = [j for j in range(ancho) if encabezado[j] or any(f[j] for f in datos)]
    columnas, vistos = [], {}
    for j in utiles:
        nombre_col = encabezado[j] or f"Columna {j + 1}"
        if nombre_col in vistos:
            vistos[nombre_col] += 1
            nombre_col = f"{nombre_col} ({vistos[nombre_col]})"
        else:
            vistos[nombre_col] = 1
        columnas.append(nombre_col)
    datos = [[f[j] for j in utiles] for f in datos]
    if not datos:
        raise ArchivoInvalido("Encontré los encabezados pero ninguna fila de datos.")
    return TablaLeida(columnas=columnas, filas=datos, fila_encabezado=numeradas[i][0], numeros=numeros,
                      advertencias=advertencias)


def perfilar(tabla: TablaLeida) -> list[dict]:
    """Tipo probable, vacíos y muestra de cada columna."""
    perfil = []
    for j, nombre in enumerate(tabla.columnas):
        valores = [f[j] for f in tabla.filas]
        llenos = [v for v in valores if v]
        numeros = sum(_es_numero(v) for v in llenos)
        fechas = sum(bool(re.match(r"^\d{1,4}[/\-.]\d{1,2}[/\-.]\d{1,4}", v)) for v in llenos)
        if llenos and fechas / len(llenos) > 0.6:
            tipo = "fecha"
        elif llenos and numeros / len(llenos) > 0.6:
            tipo = "número"
        else:
            tipo = "texto"
        perfil.append({
            "nombre": nombre, "tipo_detectado": tipo, "vacios": len(valores) - len(llenos),
            "muestra": list(dict.fromkeys(llenos))[:4],
            "texto_en_numerica": tipo == "número" and numeros < len(llenos),
        })
    return perfil
