"""Plantillas de Excel descargables para usuarios nuevos."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.ingestion.campos import campos_de

EJEMPLOS = {
    "ventas": [["15/09/2026", "Cuaderno profesional raya", 3, 45, 24, "", "Cuadernos"],
               ["15/09/2026", "Pluma BIC azul", 5, 8, 3.5, "", "Escritura"],
               ["16/09/2026", "Copias blanco y negro", 40, 1, 0.4, "", "Servicios"]],
    "productos": [["Cuaderno profesional raya", "Cuadernos", 45, 24, 40, 10, "pieza"],
                  ["Pluma BIC azul", "Escritura", 8, 3.5, 120, 30, "pieza"],
                  ["Huevo blanco", "Básicos", 48, 38, 25, 8, "kg"]],
    "compras": [["10/09/2026", "Cuaderno profesional raya", 50, 24, "", "Papelera del Centro"],
                ["12/09/2026", "Pluma BIC azul", 200, 3.5, "", "Papelera del Centro"]],
    "gastos": [["01/09/2026", "Renta del local", 1200, "Renta", "fijo"],
               ["05/09/2026", "Bolsas y empaque", 85.5, "Insumos", "variable"],
               ["30/09/2026", "Mantenimiento del mostrador", 450, "Mantenimiento", "variable"]],
}


def generar(tipo: str) -> bytes:
    campos = campos_de(tipo)
    libro = Workbook()
    hoja = libro.active
    hoja.title = tipo.capitalize()
    azul = PatternFill("solid", fgColor="5B8DB8")
    for j, campo in enumerate(campos, start=1):
        celda = hoja.cell(row=1, column=j, value=campo.etiqueta)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = azul
        celda.alignment = Alignment(horizontal="center")
        hoja.column_dimensions[celda.column_letter].width = max(16, len(campo.etiqueta) + 6)
    for fila in EJEMPLOS[tipo]:
        hoja.append(fila)

    ayuda = libro.create_sheet("Instrucciones")
    ayuda.column_dimensions["A"].width = 22
    ayuda.column_dimensions["B"].width = 14
    ayuda.column_dimensions["C"].width = 60
    ayuda.append(["Columna", "¿Obligatoria?", "Qué poner"])
    for celda in ayuda[1]:
        celda.font = Font(bold=True)
    for campo in campos:
        ayuda.append([campo.etiqueta, "Sí" if campo.requerido else "No", f"{campo.descripcion}. Ejemplo: {campo.ejemplo}"])
    ayuda.append([])
    ayuda.append(["Consejos", "", "Borra las filas de ejemplo antes de llenar. Fechas como día/mes/año. "
                                  "Montos sin texto (puedes usar $ y comas). No pasa nada si tus columnas se llaman distinto."])
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
